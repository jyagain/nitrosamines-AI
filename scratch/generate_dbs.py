import openpyxl
import json
import re
import datetime

def clean_val(val):
    if val is None:
        return None
    if isinstance(val, datetime.datetime) or isinstance(val, datetime.date):
        return val.strftime("%Y-%m-%d")
    val_str = str(val).strip()
    if val_str == "" or val_str.lower() in ["none", "nan", "null", "―", "-"]:
        return None
    return val_str

def parse_num(val):
    if val is None:
        return None
    try:
        return float(val) if '.' in str(val) else int(val)
    except:
        return None

def parse_ai(val):
    val_str = clean_val(val)
    if not val_str:
        return ""
    # Extract numeric string if possible, but keep text if NMI or special
    if "NMI" in val_str:
        return "NMI (No Mutagenic Intent)"
    cleaned = re.sub(r'[^\d.]', '', val_str)
    return cleaned if cleaned else val_str

def parse_mfds():
    wb = openpyxl.load_workbook('MFDS 1일 섭취허용량.xlsx', data_only=True)
    sheet = wb.active
    records = []
    for r in range(2, sheet.max_row + 1):
        no_val = sheet.cell(r, 1).value
        name_val = sheet.cell(r, 2).value
        if not no_val and not name_val:
            continue
        try:
            no_int = int(no_val)
        except:
            no_int = sheet.max_row - r + 1

        rec = {
            "no": no_int,
            "name": clean_val(sheet.cell(r, 2).value),
            "iupac": clean_val(sheet.cell(r, 3).value),
            "cas": clean_val(sheet.cell(r, 4).value),
            "active": clean_val(sheet.cell(r, 5).value),
            "category": parse_num(sheet.cell(r, 6).value),
            "ai": parse_ai(sheet.cell(r, 7).value),
            "remark": clean_val(sheet.cell(r, 8).value),
            "date": clean_val(sheet.cell(r, 9).value) or "2026-08",
            "publishDate": "2026년 8월 공고"
        }
        records.append(rec)
    records.sort(key=lambda x: x["no"], reverse=True)
    return records

def parse_fda():
    records = []
    current_id = 1
    # 1. CPCA
    wb_cpca = openpyxl.load_workbook('CDER Nitrosamine Impurity Acceptable Intake Limits  FDA CPCA.xlsx', data_only=True)
    sheet_cpca = wb_cpca.active
    for r in range(2, sheet_cpca.max_row + 1):
        name_val = clean_val(sheet_cpca.cell(r, 2).value)
        if not name_val:
            continue
        records.append({
            "id": current_id,
            "name": name_val,
            "api": clean_val(sheet_cpca.cell(r, 3).value),
            "category": parse_num(sheet_cpca.cell(r, 4).value),
            "ai": parse_ai(sheet_cpca.cell(r, 5).value),
            "sourceType": "CPCA",
            "publishDate": "2026년 8월 공고"
        })
        current_id += 1

    # 2. SAR
    wb_sar = openpyxl.load_workbook('CDER Nitrosamine Impurity Acceptable Intake Limits  FDA SAR.xlsx', data_only=True)
    sheet_sar = wb_sar.active
    for r in range(3, sheet_sar.max_row + 1):
        name_val = clean_val(sheet_sar.cell(r, 1).value)
        if not name_val:
            continue
        records.append({
            "id": current_id,
            "name": name_val,
            "api": clean_val(sheet_sar.cell(r, 2).value),
            "category": None,
            "ai": parse_ai(sheet_sar.cell(r, 3).value),
            "surrogate": clean_val(sheet_sar.cell(r, 4).value),
            "dateAdded": clean_val(sheet_sar.cell(r, 5).value),
            "sourceType": "SAR",
            "publishDate": "2026년 8월 공고"
        })
        current_id += 1

    # 3. Interim
    wb_interim = openpyxl.load_workbook('CDER Nitrosamine Impurity Acceptable Intake Limits  FDA interim.xlsx', data_only=True)
    sheet_interim = wb_interim.active
    for r in range(3, sheet_interim.max_row + 1):
        name_val = clean_val(sheet_interim.cell(r, 1).value)
        if not name_val:
            continue
        records.append({
            "id": current_id,
            "name": name_val,
            "api": clean_val(sheet_interim.cell(r, 2).value),
            "category": None,
            "ai": parse_ai(sheet_interim.cell(r, 3).value),
            "interimLimitPpm": clean_val(sheet_interim.cell(r, 4).value),
            "estimatedDuration": clean_val(sheet_interim.cell(r, 5).value),
            "sourceType": "Interim",
            "publishDate": "2026년 8월 공고"
        })
        current_id += 1

    return records

def parse_ema():
    records = []
    current_id = 1
    wb = openpyxl.load_workbook('Nitrosamine Impurity Acceptable Intake Limits EMA.xlsx', data_only=True)
    
    # 1. N-nitrosamines
    if 'N-nitrosamines' in wb.sheetnames:
        sheet = wb['N-nitrosamines']
        for r in range(21, sheet.max_row + 1):
            name_val = clean_val(sheet.cell(r, 2).value)
            if not name_val:
                continue
            records.append({
                "id": current_id,
                "name": name_val,
                "iupac": clean_val(sheet.cell(r, 3).value),
                "smiles": clean_val(sheet.cell(r, 4).value),
                "cas": clean_val(sheet.cell(r, 5).value),
                "acronym": clean_val(sheet.cell(r, 6).value),
                "active": clean_val(sheet.cell(r, 7).value),
                "category": parse_num(sheet.cell(r, 8).value),
                "ai": parse_ai(sheet.cell(r, 9).value),
                "remark": clean_val(sheet.cell(r, 10).value),
                "date": clean_val(sheet.cell(r, 11).value),
                "subType": "N-nitrosamines"
            })
            current_id += 1

    # 2. Other N-nitroso-structures
    if 'Other N-nitroso-structures' in wb.sheetnames:
        sheet = wb['Other N-nitroso-structures']
        for r in range(21, sheet.max_row + 1):
            name_val = clean_val(sheet.cell(r, 2).value)
            if not name_val:
                continue
            records.append({
                "id": current_id,
                "name": name_val,
                "iupac": clean_val(sheet.cell(r, 3).value),
                "smiles": clean_val(sheet.cell(r, 4).value),
                "cas": clean_val(sheet.cell(r, 5).value),
                "acronym": clean_val(sheet.cell(r, 6).value),
                "active": clean_val(sheet.cell(r, 7).value),
                "category": None,
                "ai": parse_ai(sheet.cell(r, 8).value),
                "remark": clean_val(sheet.cell(r, 9).value),
                "date": clean_val(sheet.cell(r, 10).value),
                "subType": "Other N-nitroso-structures"
            })
            current_id += 1

    return records

def parse_hc():
    records = []
    current_id = 1
    wb = openpyxl.load_workbook('Nitrosamine Impurity Acceptable Intake Limits HC.xlsx', data_only=True)
    if 'Sheet1' in wb.sheetnames:
        sheet = wb['Sheet1']
        for r in range(22, sheet.max_row + 1):
            name_val = clean_val(sheet.cell(r, 3).value)
            if not name_val:
                continue
            records.append({
                "id": current_id,
                "pubDate": clean_val(sheet.cell(r, 1).value),
                "active": clean_val(sheet.cell(r, 2).value),
                "name": name_val,
                "smiles": clean_val(sheet.cell(r, 4).value),
                "cas": clean_val(sheet.cell(r, 5).value),
                "category": parse_num(sheet.cell(r, 6).value),
                "ai": parse_ai(sheet.cell(r, 7).value)
            })
            current_id += 1
    return records

def parse_tga():
    records = []
    current_id = 1
    wb = openpyxl.load_workbook('Nitrosamine Impurity Acceptable Intake Limits TGA.xlsx', data_only=True)
    if 'Established_acceptable_intake_f' in wb.sheetnames:
        sheet = wb['Established_acceptable_intake_f']
        for r in range(2, sheet.max_row + 1):
            name_val = clean_val(sheet.cell(r, 1).value)
            if not name_val:
                continue
            records.append({
                "id": current_id,
                "name": name_val,
                "cas": clean_val(sheet.cell(r, 2).value),
                "active": clean_val(sheet.cell(r, 3).value),
                "ai": parse_ai(sheet.cell(r, 4).value),
                "category": parse_num(sheet.cell(r, 5).value),
                "pubDate": clean_val(sheet.cell(r, 6).value),
                "updatedDate": clean_val(sheet.cell(r, 7).value),
                "link": clean_val(sheet.cell(r, 8).value)
            })
            current_id += 1
    return records

if __name__ == '__main__':
    mfds_data = parse_mfds()
    fda_data = parse_fda()
    ema_data = parse_ema()
    hc_data = parse_hc()
    tga_data = parse_tga()

    print(f"Parsed MFDS: {len(mfds_data)} entries")
    print(f"Parsed FDA: {len(fda_data)} entries")
    print(f"Parsed EMA: {len(ema_data)} entries")
    print(f"Parsed HC: {len(hc_data)} entries")
    print(f"Parsed TGA: {len(tga_data)} entries")

    with open('adi_db.js', 'w', encoding='utf-8') as f:
        f.write(f"/** MFDS Official Database ({len(mfds_data)} entries) */\nwindow.MFDS_ADI_DATABASE = {json.dumps(mfds_data, indent=2, ensure_ascii=False)};\n")

    with open('fda_db.js', 'w', encoding='utf-8') as f:
        f.write(f"/** FDA Official Database ({len(fda_data)} entries) */\nwindow.FDA_ADI_DATABASE = {json.dumps(fda_data, indent=2, ensure_ascii=False)};\n")

    with open('ema_db.js', 'w', encoding='utf-8') as f:
        f.write(f"/** EMA Official Database ({len(ema_data)} entries) */\nwindow.EMA_ADI_DATABASE = {json.dumps(ema_data, indent=2, ensure_ascii=False)};\n")

    with open('hc_db.js', 'w', encoding='utf-8') as f:
        f.write(f"/** Health Canada Official Database ({len(hc_data)} entries) */\nwindow.HC_ADI_DATABASE = {json.dumps(hc_data, indent=2, ensure_ascii=False)};\n")

    with open('tga_db.js', 'w', encoding='utf-8') as f:
        f.write(f"/** TGA Official Database ({len(tga_data)} entries) */\nwindow.TGA_ADI_DATABASE = {json.dumps(tga_data, indent=2, ensure_ascii=False)};\n")

    print("All 5 JavaScript databases generated successfully!")
